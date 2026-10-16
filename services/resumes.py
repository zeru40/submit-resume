import hashlib
import uuid
from pathlib import Path
from core.store import now
from core.store import dumps
from services.profile import extract_profile

MAX_BYTES = 15 * 1024 * 1024

def import_resume(store, data_dir, filename, content):
    name = Path(filename.replace('\\', '/')).name
    suffix = Path(name).suffix.lower()
    if suffix not in {'.pdf', '.docx', '.txt'}:
        raise ValueError('支持 PDF、DOCX、TXT 简历')
    if not content or len(content) > MAX_BYTES:
        raise ValueError('简历不能为空且不能超过15MB')
    digest = hashlib.sha256(content).hexdigest()
    existing = store.one('SELECT * FROM resumes WHERE sha256=?', (digest,))
    if existing:
        return existing
    ident = uuid.uuid4().hex
    target = Path(data_dir) / 'resumes' / (ident + suffix)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(content)
    try:
        if suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(target)
            if len(reader.pages) > 30:
                raise ValueError('简历最多支持30页')
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        elif suffix == '.docx':
            from zipfile import ZipFile
            with ZipFile(target) as archive:
                if sum(i.file_size for i in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError('Word文档解压后过大')
            from docx import Document
            doc = Document(target)
            text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
        else:
            text = content.decode('utf-8-sig')
        if not text.strip():
            raise ValueError('未提取到文字；扫描版PDF请先转成可选中文字的PDF或DOCX')
        store.execute('INSERT INTO resumes(id,name,path,sha256,text,created_at,profile) VALUES(?,?,?,?,?,?,?)',
                      (ident,name,str(target),digest,text[:100000],now(),dumps(extract_profile(text))))
        return store.one('SELECT * FROM resumes WHERE id=?', (ident,))
    except BaseException as exc:
        target.unlink(missing_ok=True)
        if isinstance(exc,Exception) and not isinstance(exc,ValueError):
            raise ValueError('无法解析简历，请确认文件未损坏、未加密且格式正确') from exc
        raise
